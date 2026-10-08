"""Agente de revisión de documentos de ingeniería (pensado para correr en Kaggle).

Flujo:
  1. Carga cada documento (PDF, DOCX, XLSX/CSV, TXT/MD).
  2. Revisa cada uno contra un checklist con Claude -> hallazgos estructurados (JSON).
  3. Si hay varios documentos, hace una pasada de consistencia cruzada.
  4. Escribe un informe Markdown + hallazgos en JSON y CSV.

Uso rápido (notebook):
    from doc_review_agent import ReviewConfig, review_folder
    review_folder("/kaggle/input/mis-docs", "/kaggle/working/informe", ReviewConfig())
"""
from __future__ import annotations

import base64
import csv
import json
import os
from dataclasses import dataclass, field
from pathlib import Path

try:
    import anthropic
except ImportError:  # el backend local no necesita el SDK de Anthropic
    anthropic = None

DEFAULT_MODEL = "claude-opus-5-5"
DEFAULT_CHECKLIST = Path(__file__).with_name("checklist_default.md")
SUPPORTED = {".pdf", ".docx", ".xlsx", ".xlsm", ".csv", ".txt", ".md", ".png", ".jpg", ".jpeg", ".tif", ".tiff"}
MAX_PDF_BYTES = 30 * 1024 * 1024  # límite de la API: 32 MB por request

SEVERITIES = ["critica", "mayor", "menor", "observacion"]

SYSTEM_PROMPT = """Eres un ingeniero revisor senior que audita documentación técnica de ingeniería \
(memorias de cálculo, especificaciones, planos descritos, listas de materiales, procedimientos, \
informes de ensayo, etc.).

Reglas:
- Basa cada hallazgo en evidencia del documento: cita textualmente el fragmento o indica página/hoja/sección.
- No inventes datos. Si falta información necesaria para concluir, repórtalo como información faltante.
- Verifica aritmética, unidades, coherencia de valores entre secciones y referencias normativas citadas.
- No afirmes que una norma exige algo si no estás seguro; en ese caso indica "verificar contra la norma".
- Responde en español, con tono técnico y conciso."""

REVIEW_SCHEMA = {
    "type": "object",
    "properties": {
        "tipo_documento": {"type": "string"},
        "resumen": {"type": "string"},
        "evaluacion_global": {
            "type": "string",
            "enum": ["aprobado", "aprobado_con_comentarios", "requiere_revision", "rechazado"],
        },
        "hallazgos": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "id": {"type": "string"},
                    "severidad": {"type": "string", "enum": SEVERITIES},
                    "categoria": {"type": "string"},
                    "ubicacion": {"type": "string"},
                    "evidencia": {"type": "string"},
                    "problema": {"type": "string"},
                    "recomendacion": {"type": "string"},
                },
                "required": [
                    "id", "severidad", "categoria", "ubicacion",
                    "evidencia", "problema", "recomendacion",
                ],
                "additionalProperties": False,
            },
        },
        "informacion_faltante": {"type": "array", "items": {"type": "string"}},
    },
    "required": ["tipo_documento", "resumen", "evaluacion_global", "hallazgos", "informacion_faltante"],
    "additionalProperties": False,
}

CROSS_SCHEMA = {
    "type": "object",
    "properties": {
        "resumen": {"type": "string"},
        "inconsistencias": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "severidad": {"type": "string", "enum": SEVERITIES},
                    "documentos": {"type": "array", "items": {"type": "string"}},
                    "descripcion": {"type": "string"},
                    "recomendacion": {"type": "string"},
                },
                "required": ["severidad", "documentos", "descripcion", "recomendacion"],
                "additionalProperties": False,
            },
        },
    },
    "required": ["resumen", "inconsistencias"],
    "additionalProperties": False,
}


KAGGLE_INPUT = Path("/kaggle/input")


def is_project_file(p: Path) -> bool:
    """Archivos del propio agente (checklists, README) que no son documentos a revisar."""
    n = p.name.lower()
    return n.startswith("checklist_") or n == "readme.md" or ".git" in p.parts or "agente" in p.parts[-2:-1]


def resolve_dir(path) -> Path:
    """Acepta una carpeta, o el nombre/URL de un dataset de Kaggle (busca su carpeta en /kaggle/input)."""
    p = Path(str(path))
    if p.is_dir():
        return p
    slug = str(path).strip().rstrip("/").split("/")[-1]
    if KAGGLE_INPUT.is_dir() and slug:
        base_depth = len(KAGGLE_INPUT.parts)
        for root, dirs, _ in os.walk(KAGGLE_INPUT):
            if len(Path(root).parts) - base_depth >= 4:
                dirs[:] = []
            if Path(root).name == slug:
                return Path(root)
        found = sorted(str(Path(r).relative_to(KAGGLE_INPUT)) for r, d, _ in os.walk(KAGGLE_INPUT)
                       if len(Path(r).parts) - base_depth <= 2 and Path(r) != KAGGLE_INPUT)
        raise FileNotFoundError(
            f"No encuentro la carpeta «{path}». ¿Agregaste el dataset al notebook (Input → Add Input)? "
            f"Carpetas disponibles en /kaggle/input: {found[:20] or 'ninguna'}")
    raise FileNotFoundError(f"No existe la carpeta {path}")


@dataclass
class ReviewConfig:
    model: str = os.environ.get("DOC_REVIEW_MODEL", DEFAULT_MODEL)
    effort: str = "high"  # low | medium | high | xhigh | max
    max_tokens: int = 32000
    checklist_path: Path = DEFAULT_CHECKLIST
    extra_instructions: str = ""  # contexto del proyecto, norma aplicable, etc.
    cross_check: bool = True
    use_fallbacks: bool = True  # fallback server-side ante rechazos por clasificadores
    # --- planos (PDF impresos desde CAD / imágenes)
    plan_mode: str = "auto"  # auto: PDF de formato ≥ A3 o con nombre de plano | always | never
    plan_checklist_path: Path = Path(__file__).with_name("checklist_planos.md")
    plan_tiles: bool = True  # además de la vista general y el cajetín, revisar 4 cuadrantes ampliados
    plan_max_pages: int = 20  # tope de hojas por PDF


@dataclass
class LoadedDoc:
    name: str
    blocks: list[dict]  # bloques de contenido listos para la API
    notes: list[str] = field(default_factory=list)


# --------------------------------------------------------------------------- carga

def _text_block(text: str) -> dict:
    return {"type": "text", "text": text}


def _read_docx(path: Path) -> str:
    import docx

    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for ti, table in enumerate(d.tables, 1):
        parts.append(f"\n[Tabla {ti}]")
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def _read_xlsx(path: Path) -> str:
    import openpyxl

    # Dos pasadas: valores calculados y fórmulas, para poder revisar ambos.
    wb_v = openpyxl.load_workbook(path, data_only=True)
    wb_f = openpyxl.load_workbook(path, data_only=False)
    out = []
    for ws in wb_v.worksheets:
        wf = wb_f[ws.title]
        out.append(f"\n## Hoja: {ws.title}")
        for row_v, row_f in zip(ws.iter_rows(), wf.iter_rows()):
            cells = []
            for cv, cf in zip(row_v, row_f):
                if cv.value is None and cf.value is None:
                    continue
                if isinstance(cf.value, str) and cf.value.startswith("="):
                    cells.append(f"{cv.coordinate}={cv.value!r} (fórmula {cf.value})")
                else:
                    cells.append(f"{cv.coordinate}={cv.value!r}")
            if cells:
                out.append("; ".join(cells))
    return "\n".join(out)


def _read_pdf_text(path: Path) -> str:
    from pypdf import PdfReader

    pages = []
    for i, page in enumerate(PdfReader(str(path)).pages, 1):
        pages.append(f"[Página {i}]\n{page.extract_text() or ''}")
    return "\n\n".join(pages)


def load_document(path: str | Path, native_pdf: bool = True) -> LoadedDoc:
    """native_pdf=True envía el PDF a Claude; False extrae texto (para modelos locales)."""
    path = Path(path)
    ext = path.suffix.lower()
    if ext in {".png", ".jpg", ".jpeg", ".tif", ".tiff"}:
        raise ValueError(f"{path.name} es una imagen: se revisa como plano (plan_review) o con chat.ask_image().")
    if ext not in SUPPORTED:
        raise ValueError(f"Formato no soportado: {path.name}")

    if ext == ".pdf" and not native_pdf:
        text = _read_pdf_text(path)
        if len(text.strip()) < 50:
            raise ValueError(f"{path.name} parece un PDF escaneado (sin texto); requiere OCR o un modelo con visión.")
        return LoadedDoc(path.name, [_text_block(text)])

    if ext == ".pdf":
        data = path.read_bytes()
        if len(data) > MAX_PDF_BYTES:
            raise ValueError(f"{path.name} pesa más de 30 MB; divídelo en partes.")
        block = {
            "type": "document",
            "source": {
                "type": "base64",
                "media_type": "application/pdf",
                "data": base64.standard_b64encode(data).decode("ascii"),
            },
            "title": path.name,
        }
        return LoadedDoc(path.name, [block])

    if ext == ".docx":
        text = _read_docx(path)
    elif ext in (".xlsx", ".xlsm"):
        text = _read_xlsx(path)
    else:
        text = path.read_text(encoding="utf-8", errors="replace")

    if not text.strip():
        raise ValueError(f"{path.name} no tiene contenido de texto extraíble.")
    return LoadedDoc(path.name, [_text_block(text)])


# --------------------------------------------------------------------------- llamadas

class ClaudeLLM:
    """Backend que usa la API de Claude (requiere ANTHROPIC_API_KEY)."""

    native_pdf = True
    vision = True
    image_max_side = 2400  # px del lado mayor al renderizar planos
    chunk_chars = 0  # 0 = sin troceado (contexto de 1M tokens)

    def __init__(self, cfg: ReviewConfig, client=None):
        self.cfg = cfg
        self.client = client or anthropic.Anthropic()

    def generate_text(self, system: str, user: str, images: list[bytes] | None = None) -> str:
        content = [{"type": "image", "source": {"type": "base64", "media_type": "image/png",
                                                "data": base64.standard_b64encode(im).decode("ascii")}}
                   for im in (images or [])] + [{"type": "text", "text": user}]
        msg = self.client.messages.create(
            model=self.cfg.model, max_tokens=8000, system=system,
            output_config={"effort": "medium"},
            messages=[{"role": "user", "content": content}],
        )
        return "".join(b.text for b in msg.content if b.type == "text").strip()

    def generate_json(self, system: str, content: list[dict], schema: dict) -> dict:
        cfg, client = self.cfg, self.client
        params = dict(
            model=cfg.model,
            max_tokens=cfg.max_tokens,
            system=system,
            thinking={"type": "adaptive"},
            output_config={"effort": cfg.effort, "format": {"type": "json_schema", "schema": schema}},
            messages=[{"role": "user", "content": content}],
        )

        def run(use_fallbacks: bool):
            if use_fallbacks:
                ctx = client.beta.messages.stream(
                    **params, betas=["server-side-fallback-2026-07-01"], fallbacks="default"
                )
            else:
                ctx = client.messages.stream(**params)
            with ctx as stream:
                return stream.get_final_message()

        try:
            msg = run(cfg.use_fallbacks)
        except anthropic.BadRequestError:
            if not cfg.use_fallbacks:
                raise
            msg = run(False)  # la cuenta/plataforma no acepta fallbacks: reintenta sin ellos

        if msg.stop_reason == "refusal":
            raise RuntimeError(f"La solicitud fue rechazada por los clasificadores de seguridad: {msg.stop_details}")
        if msg.stop_reason == "max_tokens":
            raise RuntimeError("La respuesta se cortó por max_tokens; sube ReviewConfig.max_tokens.")
        text = next(b.text for b in msg.content if b.type == "text")
        return json.loads(text)


def _build_system(cfg: ReviewConfig) -> str:
    checklist = Path(cfg.checklist_path).read_text(encoding="utf-8")
    system = f"{SYSTEM_PROMPT}\n\n# Checklist de revisión\n{checklist}"
    if cfg.extra_instructions.strip():
        system += f"\n\n# Contexto adicional del proyecto\n{cfg.extra_instructions.strip()}"
    return system


def _split_text(text: str, max_chars: int) -> list[str]:
    """Trocea por líneas sin superar max_chars (una línea larga se parte)."""
    chunks, cur, size = [], [], 0
    for line in text.splitlines():
        while len(line) > max_chars:
            if cur:
                chunks.append("\n".join(cur)); cur, size = [], 0
            chunks.append(line[:max_chars]); line = line[max_chars:]
        if size + len(line) + 1 > max_chars and cur:
            chunks.append("\n".join(cur)); cur, size = [], 0
        cur.append(line); size += len(line) + 1
    if cur:
        chunks.append("\n".join(cur))
    return chunks


def _global_assessment(hallazgos: list[dict]) -> str:
    sev = {h["severidad"] for h in hallazgos}
    if "critica" in sev or "mayor" in sev:
        return "requiere_revision"
    return "aprobado_con_comentarios" if sev else "aprobado"


def review_document(llm, doc: LoadedDoc, cfg: ReviewConfig) -> dict:
    system = _build_system(cfg)
    ask = ("Revisa el documento «{name}»{part} según el checklist. "
           "Numera los hallazgos como H-01, H-02, ... ordenados por severidad.")
    limit = getattr(llm, "chunk_chars", 0)
    texts = [b["text"] for b in doc.blocks if b["type"] == "text"]
    needs_split = limit and texts and sum(len(t) for t in texts) > limit

    if not needs_split:
        content = list(doc.blocks) + [_text_block(ask.format(name=doc.name, part=""))]
        result = llm.generate_json(system, content, REVIEW_SCHEMA)
        result["documento"] = doc.name
        return result

    # Documento largo en modelo con contexto reducido: revisar por partes y fusionar.
    chunks = _split_text("\n".join(texts), limit)
    parts = []
    for i, chunk in enumerate(chunks, 1):
        print(f"  parte {i}/{len(chunks)}")
        part = f" (parte {i} de {len(chunks)}; revisa solo lo que aparece en esta parte)"
        content = [_text_block(chunk), _text_block(ask.format(name=doc.name, part=part))]
        parts.append(llm.generate_json(system, content, REVIEW_SCHEMA))

    hallazgos = [h for r in parts for h in r["hallazgos"]]
    hallazgos.sort(key=lambda h: _SEV_ORDER[h["severidad"]])
    for n, h in enumerate(hallazgos, 1):
        h["id"] = f"H-{n:02d}"
    faltante = list(dict.fromkeys(x for r in parts for x in r["informacion_faltante"]))
    return {
        "documento": doc.name,
        "tipo_documento": parts[0]["tipo_documento"],
        "resumen": " ".join(r["resumen"] for r in parts),
        "evaluacion_global": _global_assessment(hallazgos),
        "hallazgos": hallazgos,
        "informacion_faltante": faltante,
    }


def cross_check(llm, reviews: list[dict], cfg: ReviewConfig) -> dict:
    resumen = [
        {"documento": r["documento"], "tipo": r["tipo_documento"], "resumen": r["resumen"],
         **({"cajetin": r["cajetin"]} if r.get("cajetin") else {}),
         "hallazgos": [{k: h[k] for k in ("id", "severidad", "problema", "evidencia")} for h in r["hallazgos"]]}
        for r in reviews
    ]
    prompt = (
        "Estas son las revisiones individuales de varios documentos del mismo proyecto. "
        "Identifica inconsistencias ENTRE documentos (valores, unidades, revisiones, nomenclatura, "
        "alcance, referencias cruzadas). No repitas hallazgos internos de un solo documento.\n\n"
        + json.dumps(resumen, ensure_ascii=False, indent=2)
    )
    return llm.generate_json(SYSTEM_PROMPT, [_text_block(prompt)], CROSS_SCHEMA)


# --------------------------------------------------------------------------- informe

_SEV_ORDER = {s: i for i, s in enumerate(SEVERITIES)}


def render_markdown(reviews: list[dict], cross: dict | None, errors: dict[str, str]) -> str:
    out = ["# Informe de revisión de documentos de ingeniería\n"]
    total = {s: 0 for s in SEVERITIES}
    for r in reviews:
        for h in r["hallazgos"]:
            total[h["severidad"]] += 1
    out.append("## Resumen ejecutivo\n")
    out.append("| Documento | Tipo | Evaluación | Hallazgos |\n|---|---|---|---|")
    for r in reviews:
        out.append(f"| {r['documento']} | {r['tipo_documento']} | {r['evaluacion_global']} | {len(r['hallazgos'])} |")
    plans = [r for r in reviews if r.get("cajetin")]
    if plans:
        out.append("\n**Planos (datos del cajetín):**\n")
        out.append("| Hoja | N° plano | Título | Rev. | Fecha | Escala | Dibujó | Revisó | Aprobó |\n|---|---|---|---|---|---|---|---|---|")
        for r in plans:
            c = r["cajetin"]
            out.append(f"| {r['documento']} | {c['numero']} | {c['titulo']} | {c['revision']} | {c['fecha']} | "
                       f"{c['escala']} | {c['dibujo']} | {c['reviso']} | {c['aprobo']} |")
    out.append("\nTotal por severidad: " + ", ".join(f"{s}: {n}" for s, n in total.items()) + "\n")
    if errors:
        out.append("**Documentos no procesados:**\n")
        out += [f"- {n}: {e}" for n, e in errors.items()]
        out.append("")

    for r in reviews:
        out.append(f"## {r['documento']}\n")
        out.append(f"**Tipo:** {r['tipo_documento']}  \n**Evaluación:** {r['evaluacion_global']}\n")
        out.append(r["resumen"] + "\n")
        for h in sorted(r["hallazgos"], key=lambda h: _SEV_ORDER[h["severidad"]]):
            out.append(f"### {h['id']} · {h['severidad'].upper()} · {h['categoria']}")
            out.append(f"- **Ubicación:** {h['ubicacion']}")
            out.append(f"- **Evidencia:** {h['evidencia']}")
            out.append(f"- **Problema:** {h['problema']}")
            out.append(f"- **Recomendación:** {h['recomendacion']}\n")
        if r["informacion_faltante"]:
            out.append("**Información faltante:**\n")
            out += [f"- {x}" for x in r["informacion_faltante"]]
            out.append("")

    if cross:
        out.append("## Consistencia entre documentos\n")
        out.append(cross["resumen"] + "\n")
        for c in sorted(cross["inconsistencias"], key=lambda c: _SEV_ORDER[c["severidad"]]):
            out.append(f"- **[{c['severidad'].upper()}]** ({', '.join(c['documentos'])}) {c['descripcion']} "
                       f"→ _{c['recomendacion']}_")
        out.append("")
    return "\n".join(out)


def write_outputs(out_dir: Path, reviews: list[dict], cross: dict | None, errors: dict[str, str]) -> None:
    out_dir.mkdir(parents=True, exist_ok=True)
    (out_dir / "informe.md").write_text(render_markdown(reviews, cross, errors), encoding="utf-8")
    (out_dir / "revision.json").write_text(
        json.dumps({"revisiones": reviews, "consistencia": cross, "errores": errors}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    cols = ["documento", "id", "severidad", "categoria", "ubicacion", "evidencia", "problema", "recomendacion"]
    with open(out_dir / "hallazgos.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        for r in reviews:
            for h in r["hallazgos"]:
                w.writerow({"documento": r["documento"], **h})


# --------------------------------------------------------------------------- entrada

def review_folder(in_dir: str | Path, out_dir: str | Path, cfg: ReviewConfig | None = None,
                  client=None, llm=None) -> dict:
    """llm: backend con .generate_json(); por defecto Claude. Para modelo local ver local_llm.LocalLLM."""
    cfg = cfg or ReviewConfig()
    llm = llm or ClaudeLLM(cfg, client)
    from plan_review import IMAGE_EXTS, is_plan_pdf, review_plan_image, review_plan_pdf

    in_dir = resolve_dir(in_dir)
    paths = sorted(p for p in Path(in_dir).rglob("*")
                   if p.is_file() and p.suffix.lower() in SUPPORTED and not is_project_file(p))
    if not paths:
        raise FileNotFoundError(f"No hay documentos soportados en {in_dir} ({', '.join(sorted(SUPPORTED))})")

    reviews, errors = [], {}
    for p in paths:
        print(f"Revisando {p.name} ...")
        try:
            if p.suffix.lower() in IMAGE_EXTS or (p.suffix.lower() == ".pdf" and is_plan_pdf(p, cfg.plan_mode)):
                print("  (plano)")
                fn = review_plan_image if p.suffix.lower() in IMAGE_EXTS else review_plan_pdf
                reviews.extend(fn(llm, p, cfg))
                continue
            reviews.append(review_document(llm, load_document(p, llm.native_pdf), cfg))
        except Exception as e:  # un documento defectuoso no debe frenar el lote
            errors[p.name] = f"{type(e).__name__}: {e}"
            print(f"  ! {errors[p.name]}")

    cross = None
    if cfg.cross_check and len(reviews) > 1:
        print("Revisando consistencia entre documentos ...")
        try:
            cross = cross_check(llm, reviews, cfg)
        except Exception as e:
            errors["(consistencia)"] = f"{type(e).__name__}: {e}"

    write_outputs(Path(out_dir), reviews, cross, errors)
    print(f"Listo: {len(reviews)} revisados, {len(errors)} con error -> {out_dir}")
    return {"revisiones": reviews, "consistencia": cross, "errores": errors}


if __name__ == "__main__":
    import argparse

    ap = argparse.ArgumentParser(description="Revisión de documentos de ingeniería con Claude")
    ap.add_argument("entrada", help="carpeta con documentos")
    ap.add_argument("salida", help="carpeta de salida")
    ap.add_argument("--contexto", default="", help="contexto/norma aplicable")
    ap.add_argument("--effort", default="high")
    a = ap.parse_args()
    review_folder(a.entrada, a.salida, ReviewConfig(extra_instructions=a.contexto, effort=a.effort))
